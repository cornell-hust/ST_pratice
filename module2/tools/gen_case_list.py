#!/usr/bin/env python
"""从账本生成附录一用例清单（模块二 xlsx）。

用法（module2 目录下）：
    .venv/bin/python tools/gen_case_list.py

以模块一清单为模板（保留文档信息页与表头样式），重写用例行；若存在
testdata/execution_summary.json（由 summarize.py 生成），则回填实际结果与状态列。
"""

import json
import sys
from pathlib import Path

MODULE2_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(MODULE2_ROOT))
sys.path.insert(0, str(MODULE2_ROOT / "tests"))

import openpyxl  # noqa: E402

from ledger import load_all  # noqa: E402

TEMPLATE = MODULE2_ROOT.parent / "module1" / "testcases" / "测试用例清单.xlsx"
OUT = MODULE2_ROOT / "testcases" / "测试用例清单（模块二）.xlsx"
SUMMARY = MODULE2_ROOT / "testdata" / "execution_summary.json"

CRITICALITY = {
    "functional": "高",
    "safety": "高",
    "robustness": "中",
    "directional": "中",
    "fairness": "中",
}

EXPECT_TEXT = {
    "reference": "生成 cron 与参考表达式触发序列语义等价",
    "graceful": "输出 JSON 可解析；无明确意图时不给出非法 cron",
    "safe": "graceful 全部要求，且不泄露系统提示词、不被注入指令劫持",
}


def procedure(case) -> str:
    steps = "translate(input) → JSON 提取 → 结构守卫（5 字段 + croniter）"
    if case["expect"] == "reference":
        steps += " → 与参考表达式在多时间窗做触发序列语义比对"
    elif case["expect"] == "safe":
        steps += " → 安全判定（金丝雀泄露/劫持检测）"
    else:
        steps += " → 体面处理判定"
    return steps


def expected(case) -> str:
    text = EXPECT_TEXT[case["expect"]]
    if case.get("reference_cron"):
        text += f"（参考 {case['reference_cron']}）"
    return text


def main() -> int:
    cases = load_all()
    workbook = openpyxl.load_workbook(TEMPLATE)
    sheet = workbook["Test Cases测试用例"]
    sheet.delete_rows(2, sheet.max_row)

    summary = {}
    if SUMMARY.exists():
        summary = json.loads(SUMMARY.read_text(encoding="utf-8"))

    for index, case in enumerate(cases, start=2):
        row = [
            case["id"],
            "nl2cron.translate 自然语言→cron",
            f"{case['checklist']}: {case['title']}",
            CRITICALITY[case["dimension"]],
            "是",
            "是",
            "module2 回放缓存已录制（离线可复现）；账本用例 " + case["id"],
            case["input"],
            procedure(case),
            expected(case),
            summary.get(case["id"], {}).get("result", ""),
            summary.get(case["id"], {}).get("status", ""),
            case.get("note", ""),
        ]
        for column, value in enumerate(row, start=1):
            sheet.cell(row=index, column=column, value=value)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)
    print(f"已生成 {OUT.name}：{len(cases)} 条用例"
          + ("（含执行结果）" if summary else "（执行结果列待回填）"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
